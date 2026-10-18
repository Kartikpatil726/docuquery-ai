import hashlib
import io
import re
from dataclasses import dataclass
from pathlib import Path

from pypdf import PdfReader

class DocumentLoadError(Exception):
    '''Raised when a document cannot be read or has no usable text.'''

@dataclass(frozen=True)
class PageText:
    page: int
    text: str

@dataclass(frozen=True)
class LoadedDocument:
    doc_id: str
    filename: str
    pages: list[PageText]
    total_pages: int
    empty_pages: int

def compute_doc_id(path: Path) -> str:
    sha = hashlib.sha256()
    with path.open('rb') as f:
        for block in iter(lambda: f.read(1024 * 1024), b''):
            sha.update(block)
    return sha.hexdigest()[:16]

def clean_text(text: str) -> str:
    text = text.replace('\x00', '')
    text = re.sub(r'[ \t]+', ' ', text)
    text = re.sub(r' ?\n ?', '\n', text)
    text = re.sub(r'\n{3,}', '\n\n', text)
    return text.strip()

def _load_pdf(path: Path, max_pages: int) -> tuple[list[PageText], int, int]:
    reader = PdfReader(str(path))
    if reader.is_encrypted:
        raise DocumentLoadError('This PDF is password-protected.')
    total = len(reader.pages)
    if total == 0:
        raise DocumentLoadError('This PDF has no pages.')
    if total > max_pages:
        raise DocumentLoadError(f'PDF exceeds limit of {max_pages} pages.')

    pages: list[PageText] = []
    empty = 0
    for number, page in enumerate(reader.pages, start=1):
        text = clean_text(page.extract_text() or '')
        
        # Fallback: Agar page me direct text na mile toh images check karo
        if not text:
            try:
                import pytesseract
                from PIL import Image
                for img_obj in page.images:
                    img = Image.open(io.BytesIO(img_obj.data))
                    ocr_res = clean_text(pytesseract.image_to_string(img))
                    if ocr_res:
                        text += "\n" + ocr_res
            except Exception:
                pass

        if text:
            pages.append(PageText(page=number, text=text))
        else:
            empty += 1

    return pages, total, empty

def _load_docx(path: Path) -> tuple[list[PageText], int, int]:
    try:
        import docx
        doc = docx.Document(str(path))
        full_text = []
        for p in doc.paragraphs:
            if p.text.strip():
                full_text.append(p.text.strip())
        content = clean_text("\n\n".join(full_text))
        if not content:
            raise DocumentLoadError("Docx contains no usable text.")
        return [PageText(page=1, text=content)], 1, 0
    except Exception as e:
        raise DocumentLoadError(f"Failed to read DOCX: {e}")

def _load_pptx(path: Path) -> tuple[list[PageText], int, int]:
    try:
        from pptx import Presentation
        prs = Presentation(str(path))
        pages: list[PageText] = []
        empty = 0
        for idx, slide in enumerate(prs.slides, start=1):
            slide_text = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_text.append(shape.text.strip())
            t = clean_text("\n".join(slide_text))
            if t:
                pages.append(PageText(page=idx, text=t))
            else:
                empty += 1
        return pages, len(prs.slides), empty
    except Exception as e:
        raise DocumentLoadError(f"Failed to read PPTX: {e}")

def load_document(path: Path, original_filename: str, max_pages: int = 500) -> LoadedDocument:
    ext = Path(original_filename).suffix.lower()
    
    if ext == '.pdf':
        pages, total, empty = _load_pdf(path, max_pages)
    elif ext in ('.docx', '.doc'):
        pages, total, empty = _load_docx(path)
    elif ext in ('.pptx', '.ppt'):
        pages, total, empty = _load_pptx(path)
    else:
        raise DocumentLoadError(f'Unsupported format: {ext}. Only PDF, DOCX, and PPTX are supported.')

    if not pages:
        raise DocumentLoadError('No text could be extracted from this document.')

    return LoadedDocument(
        doc_id=compute_doc_id(path),
        filename=original_filename,
        pages=pages,
        total_pages=total,
        empty_pages=empty,
    )
